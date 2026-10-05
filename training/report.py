"""Publish only completed five-fold results; paired ablations use identical OOF rows."""
from pathlib import Path
import fcntl,itertools,json
import numpy as np
import pandas as pd
from scipy.stats import t as student_t
from docx import Document
from runtime import ROOT,VARIANTS
RESULT=ROOT/'training/results'
LABELS={'full':'Full Model','without_molformer':'w/o MolFormer',
 'without_descriptors':'w/o Branch Descriptors','without_cnn':'w/o CNN',
 'without_lstm':'w/o LSTM','without_augmentation':'w/o Augmentation','only_smiles':'Only SMILES Sequence'}
TARGETS=['Tg','FFV','Tc','Density','Rg']

def bootstrap(df,seed=42):
 # Resample CV groups, preserving stereo-grouped repeated rows.
 groups=df.groupby('cv_group_id').agg(n=('measured','size'),absolute_sum=('absolute_error','sum'))
 rng=np.random.default_rng(seed);a=groups.absolute_sum.to_numpy();n=groups.n.to_numpy();values=[]
 for _ in range(2000):
  ids=rng.integers(0,len(groups),len(groups));values.append(a[ids].sum()/n[ids].sum())
 return np.percentile(values,[2.5,97.5])

def holm(p):
 p=np.array(p);order=np.argsort(p);adjust=np.empty(len(p));running=0
 for rank,i in enumerate(order):
  running=max(running,(len(p)-rank)*p[i]);adjust[i]=min(1,running)
 return adjust

def md_table(df):
 rows=['| '+' | '.join(df.columns)+' |','| '+' | '.join(['---']*len(df.columns))+' |']
 for row in df.itertuples(index=False,name=None):rows.append('| '+' | '.join(map(str,row))+' |')
 return '\n'.join(rows)

def main():
 RESULT.mkdir(exist_ok=True)
 lock=(RESULT/'.report.lock').open('w')
 try:fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
 except BlockingIOError:
  print('Another report writer is active; skipping this invocation.');return
 rows=[];oof={};folds={};comparisons=[]
 for variant,target in itertools.product(VARIANTS,TARGETS):
  files=sorted((RESULT/variant/target).glob('fold_*/metrics.json'))
  if len(files)!=5:continue
  metrics=[json.loads(f.read_text()) for f in files]
  if {m['fold'] for m in metrics}!={1,2,3,4,5}:continue
  frames=[pd.read_csv(f.parent/'oof_predictions.csv') for f in files];df=pd.concat(frames,ignore_index=True)
  expected=pd.read_csv(ROOT/'training/artifacts/outer_folds.csv')
  expected=expected[expected.target==target]
  assert set(df.row_id)==set(expected.row_id) and not df.row_id.duplicated().any()
  df['absolute_error']=abs(df.measured-df.predicted);df['squared_error']=(df.measured-df.predicted)**2
  low,high=bootstrap(df)
  vals=np.array([m['mae'] for m in metrics]);oof[(variant,target)]=df;folds[(variant,target)]=metrics
  row={'variant':variant,'target':target,'n':len(df),'oof_mae':df.absolute_error.mean(),
   'fold_mae_mean':vals.mean(),'fold_mae_sd':vals.std(ddof=1),
   'oof_rmse':np.sqrt(df.squared_error.mean()),'oof_r2':1-df.squared_error.sum()/((df.measured-df.measured.mean())**2).sum(),
   'mae_ci95_low':low,'mae_ci95_high':high}
  rows.append(row);df.to_csv(RESULT/f'OOF_{variant}_{target}.csv',index=False)
 if not rows:
  print('No five-fold target/variant complete yet; no performance claims.');return
 scores=pd.DataFrame(rows);scores.to_csv(RESULT/'performance_summary.csv',index=False)
 for variant,target in itertools.product(list(VARIANTS)[1:],TARGETS):
  if (variant,target) not in oof or ('full',target) not in oof:continue
  a=oof[(variant,target)].set_index('row_id');b=oof[('full',target)].set_index('row_id')
  assert a.index.sort_values().equals(b.index.sort_values())
  assert a.fold.sort_index().equals(b.fold.sort_index())
  delta=a.absolute_error.sort_index()-b.absolute_error.sort_index()
  d=[];ratios=[]
  for fold in range(1,6):
   mask=a.fold.sort_index()==fold;d.append(delta[mask].mean())
   m=next(m for m in folds[(variant,target)] if m['fold']==fold);ratios.append(m['n_test']/m['n_train'])
  d=np.array(d);se=np.sqrt((1/5+np.mean(ratios))*d.var(ddof=1))
  p=1. if se==0 and d.mean()==0 else (0. if se==0 else float(2*student_t.sf(abs(d.mean()/se),df=4)))
  comparisons.append({'variant':variant,'target':target,'delta_oof_mae_vs_full':delta.mean(),
   'fold_delta_mean':d.mean(),'fold_delta_sd':d.std(ddof=1),'corrected_resampled_t_p':p,
   'test_train_ratio_mean':np.mean(ratios)})
 if comparisons:
  comparisons=pd.DataFrame(comparisons);comparisons['holm_p']=holm(comparisons.corrected_resampled_t_p)
  comparisons['holm_family_size']=len(comparisons)
  comparisons['multiple_test_scope']='all currently complete comparisons; final 30 tests only after all complete'
  comparisons.to_csv(RESULT/'ablation_comparisons.csv',index=False)
 display=[]
 for variant in VARIANTS:
  r={'Model Variant':LABELS[variant]}
  for target in TARGETS:
   m=scores[(scores.variant==variant)&(scores.target==target)]
   digits=2 if target=='Tg' else (3 if target=='Rg' else 5)
   r[target]='pending' if m.empty else f"{m.iloc[0].fold_mae_mean:.{digits}f} ± {m.iloc[0].fold_mae_sd:.{digits}f}"
  display.append(r)
 table=pd.DataFrame(display);table.to_csv(RESULT/'Table4_ablation.csv',index=False)
 detailed=scores.copy()
 if len(comparisons):
  detailed=detailed.merge(comparisons[['variant','target','delta_oof_mae_vs_full',
      'corrected_resampled_t_p','holm_p','holm_family_size']],on=['variant','target'],how='left')
 detailed.to_csv(RESULT/'Table4_detailed.csv',index=False)
 md=['# Actual retraining and ablation results','',md_table(table),'',
  'Values are outer five-fold MAE mean ± sample SD, in original units. Each filled cell represents five completed, newly trained models. Pending cells are not results.',
  'Variants share outer folds and inner early-stopping rows. Selection uses a training-only inner holdout; models are refit on all outer-training rows for the selected number of epochs. Outer validation never selects epochs.',
  'See performance_summary.csv for OOF MAE, RMSE, R² and 95% group bootstrap confidence intervals. Ablation p-values use paired five-fold MAE differences with corrected resampled variance (1/5 + mean n_test/n_train) and Holm correction across available comparisons. With only five folds, inferential conclusions remain limited.',
  'Architecture is the user-specified original inference CNN/BiLSTM, using its fixed 64-dimensional MolFormer extraction. Batch size, fold-local preprocessing, grouping and nested early stopping were updated and documented in experiment_config.json. These are new experiments, not exact numerical reproduction of the paper.',
  'The original descriptor function declares 17 fields, but only six are nonzero in these data. The descriptor ablation measures the contribution of those six active original fields; see artifacts/descriptor_audit.csv.',
  'Grouping combines canonical non-isomeric SMILES; general polymer repeat-unit equivalence is not claimed. Supplemental labels use the author training data policy from Table 1. Frozen MolFormer features have no descriptor fallback. Only training SMILES are augmented.']
 (RESULT/'Table4_ablation.md').write_text('\n'.join(md)+'\n')
 with pd.ExcelWriter(RESULT/'retraining_results.xlsx',engine='openpyxl') as writer:
  table.to_excel(writer,sheet_name='Table4',index=False);scores.to_excel(writer,sheet_name='Performance',index=False)
  detailed.to_excel(writer,sheet_name='Table4_detailed',index=False)
  if len(comparisons):comparisons.to_excel(writer,sheet_name='Paired_comparisons',index=False)
 doc=Document();doc.add_paragraph('Table 4. Actual retraining ablation: five-fold MAE mean ± SD')
 tbl=doc.add_table(rows=1,cols=len(table.columns));tbl.style='Table Grid'
 for i,c in enumerate(table.columns):tbl.rows[0].cells[i].text=c
 for row in table.itertuples(index=False,name=None):
  cells=tbl.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 doc.add_paragraph('Pending entries denote incomplete experiments. See experiment_config.json and Table4_ablation.md for validation protocol and changes from original code.')
 doc.save(RESULT/'Table4_ablation.docx')
 # Parity plots for the full model only, using genuine outer-fold predictions.
 import matplotlib;matplotlib.use('Agg');import matplotlib.pyplot as plt
 for target in TARGETS:
  if ('full',target) not in oof:continue
  d=oof[('full',target)];fig,ax=plt.subplots(figsize=(4,4))
  ax.scatter(d.measured,d.predicted,s=7,alpha=.5);lo=min(d.measured.min(),d.predicted.min());hi=max(d.measured.max(),d.predicted.max())
  ax.plot([lo,hi],[lo,hi],'k--');ax.set_xlabel('Observed');ax.set_ylabel('Out-of-fold prediction');ax.set_title(target)
  fig.tight_layout();fig.savefig(RESULT/f'parity_{target}.png',dpi=250);plt.close(fig)
 print(table.to_string(index=False))
if __name__=='__main__':main()
