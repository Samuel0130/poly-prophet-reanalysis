#!/usr/bin/env python3
"""Rebuild reviewer-requested dataset Table 1; never train or run a whole notebook.
Uses only two audited pure data functions extracted from the author's training
notebook, and verifies them against the supplied original inference canonicalizer.
"""
from pathlib import Path
from datetime import datetime, timezone
import ast, hashlib, itertools, json, platform
import numpy as np
import pandas as pd
from rdkit import Chem, rdBase

ROOT=Path(__file__).resolve().parent
DATA=ROOT/'kaggle_inputs/neurips-open-polymer-prediction-2025'
OUT=ROOT/'table1_rebuilt'
TRAIN_NOTEBOOK=ROOT/'kaggle_source/original_training/cnnlstm-bamod4c/cnnlstm-bamod4c.ipynb'
ORIGINAL=ROOT/'cnnlstm-bamod4c-molfold-aug-infer.ipynb'
TARGETS=['Tg','FFV','Tc','Density','Rg']
UNITS={'Tg':'°C','FFV':'dimensionless','Tc':'W/(m·K)','Density':'g/cm³','Rg':'Å'}
FILES={'train':'train.csv','dataset1':'train_supplement/dataset1.csv',
 'dataset2':'train_supplement/dataset2.csv','dataset3':'train_supplement/dataset3.csv',
 'dataset4':'train_supplement/dataset4.csv','test':'test.csv'}
SUPPLEMENTS=[('dataset1','Tc'),('dataset3','Tg'),('dataset4','FFV')]


def notebook_functions(path,names):
 n=json.loads(path.read_text(encoding='utf-8'))
 source='\n'.join(''.join(c.get('source',[])) for c in n['cells'][1:])
 tree=ast.parse(source)
 selected=[node for node in tree.body if isinstance(node,ast.FunctionDef) and node.name in names]
 if {node.name for node in selected}!=set(names): raise ValueError('Notebook data functions missing')
 namespace={'pd':pd,'Chem':Chem}
 exec(compile(ast.Module(body=selected,type_ignores=[]),str(path),'exec'),namespace)
 excerpts='\n\n'.join(ast.get_source_segment(source,node) for node in selected)
 return namespace,excerpts


def gid(smiles):
 return hashlib.sha256(smiles.encode()).hexdigest() if isinstance(smiles,str) else None


def stats(df,label):
 rows=[]
 for t in TARGETS:
  values=df[t].dropna();count=len(values)
  rows.append({'dataset':label,'property':t,'unit':UNITS[t],
   'unique_polymers_total':len(df),'labeled_polymers':count,'missing_count':len(df)-count,
   'missing_percent':100*(len(df)-count)/len(df),'minimum':float(values.min()),
   'maximum':float(values.max()),'mean':float(values.mean()),'sample_sd':float(values.std(ddof=1))})
 return pd.DataFrame(rows)


def md_table(df):
 columns=list(df.columns)
 lines=['| '+' | '.join(columns)+' |','| '+' | '.join(['---']*len(columns))+' |']
 for row in df.itertuples(index=False,name=None):
  lines.append('| '+' | '.join(str(v) for v in row)+' |')
 return '\n'.join(lines)


def display_stats(df):
 rows=[]
 for r in df.to_dict('records'):
  t=r['property'];digits={'Tg':2,'FFV':5,'Tc':5,'Density':5,'Rg':3}[t]
  fmt=lambda x:f'{x:.{digits}f}'
  rows.append({'Property':t,'Unit':r['unit'],'Labeled polymers':r['labeled_polymers'],
   'Range':fmt(r['minimum'])+'–'+fmt(r['maximum']),
   'Mean ± SD':fmt(r['mean'])+' ± '+fmt(r['sample_sd']),
   'Missing, n (%)':f"{r['missing_count']:,} ({r['missing_percent']:.2f}%)"})
 return pd.DataFrame(rows)


def main():
 OUT.mkdir(exist_ok=True)
 ns,excerpts=notebook_functions(TRAIN_NOTEBOOK,['clean_and_validate_smiles','add_extra_data'])
 original_ns,original_excerpt=notebook_functions(ORIGINAL,['clean_and_validate_smiles'])
 canonicalize=original_ns['clean_and_validate_smiles']
 (OUT/'audited_notebook_functions.py').write_text('from rdkit import Chem\nimport pandas as pd\n\n'+excerpts+'\n',encoding='utf-8')
 frames={};summary=[];invalid=[];dups=[];raw_stats=[]
 for source,rel in FILES.items():
  df=pd.read_csv(DATA/rel).rename(columns={'TC_mean':'Tc'})
  df.insert(0,'source',source);df.insert(1,'source_row',np.arange(len(df)))
  df['canonical_SMILES']=df['SMILES'].map(canonicalize)
  # Compare all sources, not just a few example strings.
  trained_canon=df['SMILES'].map(ns['clean_and_validate_smiles'])
  if not df['canonical_SMILES'].fillna('INVALID').equals(trained_canon.fillna('INVALID')):
   raise AssertionError('Original inference and training canonicalization differ')
  df['polymer_group_id']=df['canonical_SMILES'].map(gid)
  for t in TARGETS:
   if t not in df: df[t]=np.nan
   df[t]=pd.to_numeric(df[t],errors='raise')
   if not np.isfinite(df[t].dropna()).all(): raise ValueError(f'{source}: nonfinite {t}')
  good=df[df['canonical_SMILES'].notna()]
  rec={'source':source,'raw_rows':len(df),'valid_rows':len(good),
       'invalid_rows':len(df)-len(good),'unique_canonical_polymers':good['canonical_SMILES'].nunique(),
       'duplicate_valid_rows':len(good)-good['canonical_SMILES'].nunique()}
  for t in TARGETS:
   rec[t+'_labels_raw']=int(good[t].notna().sum())
   rec[t+'_labeled_unique']=int(good.loc[good[t].notna(),'canonical_SMILES'].nunique())
  summary.append(rec);frames[source]=df
  invalid.extend(df[df['canonical_SMILES'].isna()].to_dict('records'))
  for smiles,group in good.groupby('canonical_SMILES'):
   if len(group)<2: continue
   for t in TARGETS:
    vals=group[t].dropna()
    if len(vals):
     dups.append({'source':source,'canonical_SMILES':smiles,'polymer_group_id':gid(smiles),
      'property':t,'rows':len(group),'labeled_rows':len(vals),'distinct_labels':vals.nunique(),
      'minimum':vals.min(),'maximum':vals.max(),'mean':vals.mean(),
      'numerical_disagreement':not np.isclose(vals.min(),vals.max(),rtol=1e-6,atol=1e-8),
      'source_rows':';'.join(str(x) for x in group['source_row'])})
  raw_stats.append(stats(good.assign(SMILES=good['canonical_SMILES']),source))
  print(source,rec,flush=True)
 pd.DataFrame(summary).to_csv(OUT/'source_summary.csv',index=False)
 pd.concat(frames.values(),ignore_index=True).to_csv(OUT/'canonicalized_source_rows.csv',index=False)
 pd.DataFrame(invalid,columns=list(frames['train'].columns)).to_csv(OUT/'invalid_smiles.csv',index=False)
 pd.DataFrame(dups,columns=['source','canonical_SMILES','polymer_group_id','property','rows',
  'labeled_rows','distinct_labels','minimum','maximum','mean','numerical_disagreement','source_rows']).to_csv(OUT/'within_source_duplicates.csv',index=False)
 pd.concat(raw_stats,ignore_index=True).rename(columns={'unique_polymers_total':'source_valid_rows',
  'labeled_polymers':'labeled_rows'}).to_csv(OUT/'raw_source_property_statistics.csv',index=False)

 sets={s:set(d['canonical_SMILES'].dropna()) for s,d in frames.items()}
 overlap=[]
 for left,right in itertools.combinations(frames,2):
  overlap.append({'source_a':left,'source_b':right,'overlap_unique':len(sets[left]&sets[right]),
   'a_only_unique':len(sets[left]-sets[right]),'b_only_unique':len(sets[right]-sets[left])})
 pd.DataFrame(overlap).to_csv(OUT/'source_overlap.csv',index=False)
 conflicts=[]
 for left,right in itertools.combinations(frames,2):
  a=frames[left].dropna(subset=['canonical_SMILES']).groupby('canonical_SMILES')[TARGETS].mean()
  b=frames[right].dropna(subset=['canonical_SMILES']).groupby('canonical_SMILES')[TARGETS].mean()
  for t in TARGETS:
   pair=a[[t]].join(b[[t]],how='inner',lsuffix='_a',rsuffix='_b').dropna()
   for smiles,r in pair.iterrows():
    va,vb=float(r[t+'_a']),float(r[t+'_b'])
    conflicts.append({'source_a':left,'source_b':right,'property':t,
     'canonical_SMILES':smiles,'polymer_group_id':gid(smiles),'value_a':va,'value_b':vb,
     'absolute_difference':abs(va-vb),'numerical_disagreement':not bool(np.isclose(va,vb,rtol=1e-6,atol=1e-8)),
     'resolution':'retain train label' if left=='train' else 'report only'})
 conflict_df=pd.DataFrame(conflicts,columns=['source_a','source_b','property','canonical_SMILES',
  'polymer_group_id','value_a','value_b','absolute_difference','numerical_disagreement','resolution'])
 conflict_df.to_csv(OUT/'overlapping_label_audit.csv',index=False)
 conflict_df[conflict_df['numerical_disagreement']==True].to_csv(OUT/'label_conflicts.csv',index=False)

 # Execute the author's exact small merge function, never the training notebook.
 primary=frames['train'].dropna(subset=['canonical_SMILES'])[['canonical_SMILES']+TARGETS].rename(columns={'canonical_SMILES':'SMILES'}).copy()
 if primary['SMILES'].duplicated().any(): raise ValueError('Primary duplicate policy needs review')
 merged=primary.copy();merge_steps=[]
 label_origin={t:dict(zip(primary.loc[primary[t].notna(),'SMILES'],['train']*int(primary[t].notna().sum()))) for t in TARGETS}
 supplement_means={}
 for source,t in SUPPLEMENTS:
  extra=frames[source][['SMILES',t]].copy()
  grouped=frames[source].dropna(subset=['canonical_SMILES',t]).groupby('canonical_SMILES')[t].mean()
  supplement_means[t]=grouped
  prior=set(merged['SMILES']);old_count=int(merged[t].notna().sum());old_size=len(merged)
  merged=ns['add_extra_data'](merged,extra,t)
  for smiles in grouped.index:
   if smiles not in label_origin[t]:label_origin[t][smiles]=source
  merge_steps.append({'source':source,'property':t,'raw_rows':len(extra),'unique_labeled_structures':len(grouped),
   'overlap_with_current_structures':len(set(grouped.index)&prior),'new_structures':len(merged)-old_size,
   'new_labels':int(merged[t].notna().sum())-old_count,'unique_structures_after':len(merged)})
 if merged['SMILES'].isna().any() or merged['SMILES'].duplicated().any():raise AssertionError('Merged identity uniqueness failed')
 if set(merged['SMILES']) != sets['train']|sets['dataset1']|sets['dataset3']|sets['dataset4']:
  raise AssertionError('Merged membership mismatch')
 if not merged[TARGETS].notna().any(axis=1).all():raise AssertionError('Unlabeled row in merged analysis set')
 for t in TARGETS:
  # Primary labels must survive the merge exactly; every added label is a supplement mean.
  check=merged.set_index('SMILES')[t];orig=primary.set_index('SMILES')[t].dropna()
  if not np.allclose(check.loc[orig.index],orig,rtol=0,atol=0):raise AssertionError('Primary labels overwritten')
  origins=merged['SMILES'].map(label_origin[t]);merged[t+'_label_source']=origins
  if not origins.notna().equals(merged[t].notna()):raise AssertionError('Missing label provenance')
  if t in supplement_means:
   keys=origins.index[origins.notna() & origins.ne('train')]
   smiles=merged.loc[keys,'SMILES']
   if not np.allclose(merged.loc[keys,t],supplement_means[t].loc[smiles].to_numpy(),rtol=0,atol=1e-12):
    raise AssertionError('Supplement means mismatch')
 merged.insert(0,'polymer_group_id',merged['SMILES'].map(gid))
 merged['structure_sources']=merged['SMILES'].map(lambda s:';'.join(name for name in ['train','dataset1','dataset3','dataset4'] if s in sets[name]))
 merged.to_csv(OUT/'merged_labeled_polymers.csv',index=False)
 pd.DataFrame(merge_steps).to_csv(OUT/'merge_steps.csv',index=False)
 label_provenance=[]
 for t in TARGETS:
  for s,n in merged[t+'_label_source'].dropna().value_counts().items():
   label_provenance.append({'property':t,'retained_label_source':s,'label_count':n})
 pd.DataFrame(label_provenance).to_csv(OUT/'retained_label_provenance.csv',index=False)
 pattern=merged[TARGETS].notna().apply(lambda r:';'.join(t for t in TARGETS if r[t]),axis=1).value_counts().rename_axis('available_targets').reset_index(name='unique_polymers')
 pattern.to_csv(OUT/'label_completeness.csv',index=False)
 primary_stats=stats(primary,'primary_train');final_stats=stats(merged,'reconstructed_merged')
 primary_stats.to_csv(OUT/'Table1_primary_train.csv',index=False)
 final_stats.to_csv(OUT/'Table1.csv',index=False,float_format='%.12g')
 shown=display_stats(final_stats);shown.to_csv(OUT/'Table1_formatted.csv',index=False)

 all_train_union=sets['train']|sets['dataset1']|sets['dataset2']|sets['dataset3']|sets['dataset4']
 train_labeled=set(merged['SMILES'])
 provenance={
  'createdUTC':datetime.now(timezone.utc).isoformat(),'python':platform.python_version(),
  'pandas':pd.__version__,'numpy':np.__version__,'rdkit':rdBase.rdkitVersion,
  'original_inference_notebook':str(ORIGINAL.relative_to(ROOT)),
  'author_training_notebook':str(TRAIN_NOTEBOOK.relative_to(ROOT)),
  'training_kernel_ref':'bfladderbean/cnnlstm-bamod4c','training_kernel_version':1,
  'training_source_sha256':hashlib.sha256(TRAIN_NOTEBOOK.read_bytes()).hexdigest(),
  'original_inference_sha256':hashlib.sha256(ORIGINAL.read_bytes()).hexdigest(),
  'identity':'RDKit canonical isomeric SMILES, exact original inference canonicalizer',
  'sd':'sample standard deviation, ddof=1, nonmissing unscaled labels only',
  'merge_policy':'author training function: supplemental duplicate mean, preserve primary labels, fill missing or append new identity',
  'historical_checkpoint_training_membership':'not independently established: output has weights only, no row/fold manifest',
  'all_downloaded_train_source_unique_structures':len(all_train_union),
  'merged_labeled_unique_polymers':len(merged),
  'dataset2_unique_unlabeled':len(sets['dataset2']),
  'dataset2_overlap_with_merged':len(sets['dataset2']&train_labeled),
  'dataset2_only_unique':len(sets['dataset2']-train_labeled),
  'visible_test_unique':len(sets['test']),'visible_test_overlap_with_merged':len(sets['test']&train_labeled),
  'label_conflict_count':int(conflict_df['numerical_disagreement'].sum()),
  'labels':dict(zip(final_stats['property'],final_stats['labeled_polymers'])),
  'verification':'canonicalizers agree for every input; unique merged identities; exact membership union; primary labels unchanged; supplement means and label provenance checked'}
 (OUT/'audit_metadata.json').write_text(json.dumps(provenance,ensure_ascii=False,indent=2),encoding='utf-8')
 notes=[
  f"The analysis dataset contains {len(merged):,} unique canonical polymer SMILES after merging train.csv with labeled supplements dataset1, dataset3 and dataset4. Counts exclude SMILES augmentation, test data, model weights and descriptor/install files.",
  'Supplemental duplicate structures are averaged within each property. Existing primary labels are retained; supplements fill missing primary labels or add new structures. No pseudo-label source is included.',
  'Labeled polymers are counted separately for each property. Missing percentages use the same merged dataset size as denominator. Mean and sample SD (ddof = 1) use only observed labels in original units, before target scaling.',
  'Tg is reported in °C, following the competition documentation. Canonical SMILES preserve specified stereochemistry; this identity rule does not establish equivalence under all polymer repeat-unit shifts or periodic representations.',
  'This table is recomputed from downloaded data using the author training notebook v1 and the supplied original inference canonicalizer. It is not independent proof of the exact membership of the historical 25 checkpoints; their output contains weights without row/fold manifests.',
 ]
 md=['# Table 1. Composition and target distributions of the reconstructed polymer analysis dataset',
  '',f'**Unique canonical polymers: {len(merged):,}.**','',md_table(shown),'','Notes:']
 md += [f'{i}. {note}' for i,note in enumerate(notes,1)]
 (OUT/'Table1.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
 source_table=pd.DataFrame(summary)[['source','raw_rows','valid_rows','unique_canonical_polymers','duplicate_valid_rows']+[t+'_labels_raw' for t in TARGETS]]
 report=['# 按审稿建议重制 Table 1：数据核验说明','',
  '## 代码版本依据','',
  '以用户指定的 cnnlstm-bamod4c-molfold-aug-infer.ipynb 为原始推理版本。该文件只读取 test.csv；它的 25 个折模型已全部下载。',
  '另外通过 Kaggle 官方 API 找到作者的 bfladderbean/cnnlstm-bamod4c v1 训练版，只提取并运行其中两个数据函数；没有执行整份 Notebook，也没有使用 fork 作为本表的数据处理依据。',
  '训练版可复现规范化和补充数据合并规则，但其与 2025 年保存的模型权重之间没有逐行训练清单可核对。因此本文 Table 1 明确标为重建分析数据集，不能声称已验证历史模型的全部训练样本。','',
  '## Table 1','',f'合并后共有 **{len(merged):,} 个规范化 SMILES 结构**。','',md_table(shown),'',
  '## 原始来源与去重','',md_table(source_table),'',
  'source_summary.csv 的标签数为原始有效行数；Table1.csv 的标签数为最终唯一结构数。两者不能直接相加。','',
  '## 合并步骤','',md_table(pd.DataFrame(merge_steps)),'',
  '同一补充来源的重复结构先取标签均值；主数据已有标签优先保留，不覆盖；其余标签补缺或新增结构。每个保留标签的来源记录在 merged_labeled_polymers.csv。','',
  '## 重叠、冲突和无标签数据','',
  f"跨来源同结构同性质的数值不一致共 {provenance['label_conflict_count']} 组，完整差异见 label_conflicts.csv；判断阈值 rtol=1e-6、atol=1e-8，仅表示数值差异，不代表超出实验/模拟误差。来源内重复见 within_source_duplicates.csv。",
  f"dataset2 有 {len(sets['dataset2']):,} 个唯一结构，但下载文件没有任何五性质标签；与合并集重叠 {len(sets['dataset2']&train_labeled):,} 个，额外无标签结构 {len(sets['dataset2']-train_labeled):,} 个，不进入 Table 1 标签统计。",
  f"全部五份训练来源的结构并集为 {len(all_train_union):,}；其中未标记结构不得当作带标签训练样本。可见 test.csv 仅 3 行，与本次合并集重叠 {len(sets['test']&train_labeled)} 个；这不代表隐藏测试集的重叠情况。",
  'sample_submission.csv 的零值是提交格式占位值，不能统计为标签。PI1M 在 fork 中虽有路径，但不是下载源，也不是本次原始训练版的数据源。','',
  '## 与审稿意见的对应及论文使用','',
  '审稿意见整合 PDF 第 25 页规定 Table 1 为最终数据组成：unique polymers、各性质标签数、范围、mean/std、缺失值。本表已包含这些字段；原论文 Table 1 的消融性能应另移至审稿建议的 Table 4，不能把本次统计解释为重新训练的性能。',
  'Tg 的绝对标签按竞赛说明为 °C；原论文曾写 K，论文中的绝对温度分布应统一单位。温差 MAE 的 °C 与 K 数值相同，但这不能用于解释原文 53.39 与 4.02 等不同 MAE。',
  '规范化 SMILES 去重保持立体化学。尚未完成周期重复单元、连接点平移等聚合物等价分组，不能据本表宣称已解决审稿人提出的全部交叉验证泄漏问题；训练与 grouped CV 属于后续性能表工作。',
  '数据增强表示不计为独立聚合物；均值/标准差按非缺失原始单位标签计算，SD 使用 ddof=1。','',
  '## 可复现运行','',
  '在项目目录执行：','.venv/bin/python download_kaggle_inputs.py --verify-only',
  '.venv/bin/python rebuild_table1.py','',
  'Table1.docx 可直接复制表格，Table1.tex 可用于 LaTeX；Table1.xlsx 包含来源、缺失、标签来源及冲突概要。保留原始 CSV、PDF 和两个用户提供的 Notebook 不变。']
 (OUT/'数据核验说明.md').write_text('\n'.join(report)+'\n',encoding='utf-8')

 # Editable manuscript formats.
 from docx import Document
 from docx.shared import Inches,Pt
 doc=Document();sec=doc.sections[0];sec.page_width=Inches(11.7);sec.page_height=Inches(8.3)
 sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.6)
 doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(10)
 doc.add_paragraph('Table 1. Composition and target distributions of the reconstructed polymer analysis dataset.')
 doc.add_paragraph(f'Unique canonical polymers: {len(merged):,}.')
 table=doc.add_table(rows=1,cols=len(shown.columns));table.style='Table Grid'
 for i,col in enumerate(shown.columns):table.rows[0].cells[i].text=col
 for row in shown.itertuples(index=False,name=None):
  cells=table.add_row().cells
  for i,value in enumerate(row):cells[i].text=str(value)
 for i,note in enumerate(notes,1):doc.add_paragraph(f'{i}. {note}')
 doc.save(OUT/'Table1.docx')
 def latex_escape(s):
  s=str(s).replace('dimensionless','1').replace('°C',r'$^\circ$C').replace('Å',r'$\mathrm{\AA}$').replace('g/cm³',r'g/cm$^3$')
  return s.replace('±',r'$\pm$').replace('–','--').replace('%',r'\%').replace('·',r'$\cdot$')
 tex=[r'\begin{table}[htbp]',r'\centering',r'\caption{Composition and target distributions of the reconstructed polymer analysis dataset.}',r'\label{tab:dataset-composition}',r'\begin{tabular}{llrllr}',r'\hline']
 tex.append(' & '.join(latex_escape(c) for c in shown.columns)+r' \\');tex.append(r'\hline')
 for row in shown.itertuples(index=False,name=None):tex.append(' & '.join(latex_escape(v) for v in row)+r' \\')
 tex += [r'\hline',r'\end{tabular}',r'\begin{minipage}{\linewidth}',r'\footnotesize',f'Unique canonical polymers: {len(merged):,}.']
 for i,note in enumerate(notes,1):tex.append(latex_escape(f'({i}) {note}')+r'\par')
 tex += [r'\end{minipage}',r'\end{table}']
 (OUT/'Table1.tex').write_text('\n'.join(tex)+'\n',encoding='utf-8')
 with pd.ExcelWriter(OUT/'Table1.xlsx',engine='openpyxl') as writer:
  for sheet,frame in [('Table1',shown),('Exact_statistics',final_stats),('Primary_train',primary_stats),
    ('Sources',pd.DataFrame(summary)),('Merge_steps',pd.DataFrame(merge_steps)),('Label_provenance',pd.DataFrame(label_provenance)),
    ('Source_overlap',pd.DataFrame(overlap)),('Label_conflicts',conflict_df[conflict_df['numerical_disagreement']==True]),
    ('Within_source_duplicates',pd.DataFrame(dups)),('Label_completeness',pattern),('Notes',pd.DataFrame({'notes':notes}))]:
   frame.to_excel(writer,sheet_name=sheet,index=False)
  for ws in writer.book.worksheets:
   ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
   for col in ws.columns:
    letter=col[0].column_letter
    ws.column_dimensions[letter].width=min(65,max(12,max(len(str(c.value or '')) for c in col)+2))
 print('\nTable 1:',flush=True);print(shown.to_string(index=False),flush=True)
 print('PASS:',provenance['verification'],flush=True)

if __name__=='__main__':main()
