"""Audit shared OOF rows and report paired differences with group bootstrap CIs."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
from docx import Document
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'external_baselines/results'

def main():
    audits=[];comparisons=[]
    for method in ['chen_rnn','pugraph']:
        for target in ['Tg','FFV','Tc','Density','Rg']:
            p=OUT/f'OOF_{method}_{target}.csv'
            if not p.exists():continue
            m=pd.read_csv(ROOT/f'training/results/OOF_full_{target}.csv')
            b=pd.read_csv(p)
            assert not m.row_id.duplicated().any() and not b.row_id.duplicated().any()
            assert set(m.row_id)==set(b.row_id)
            d=m.merge(b,on='row_id',suffixes=('_main','_baseline'),validate='one_to_one')
            assert np.array_equal(d.fold_main,d.fold_baseline)
            assert np.allclose(d.measured_main,d.measured_baseline,rtol=1e-9,atol=1e-8)
            assert np.isfinite(d.predicted_main).all() and np.isfinite(d.predicted_baseline).all()
            delta=d.absolute_error_baseline-d.absolute_error_main
            # Positive difference favors our full model. Resample entire groups.
            g=pd.DataFrame({'group':d.cv_group_id_main,'delta':delta}).groupby('group').agg(total=('delta','sum'),n=('delta','size'))
            rng=np.random.default_rng(42);boot=[]
            for _ in range(2000):
                ids=rng.integers(0,len(g),size=len(g));z=g.iloc[ids];boot.append(z.total.sum()/z.n.sum())
            lo,hi=np.quantile(boot,[.025,.975])
            fold_delta=d.assign(delta=delta).groupby('fold_main').delta.mean()
            audits.append(dict(method=method,target=target,n=len(d),groups=len(g),same_rows=True,same_folds=True,same_labels=True,finite=True))
            comparisons.append(dict(method=method,target=target,n=len(d),main_oof_mae=d.absolute_error_main.mean(),
                                    baseline_oof_mae=d.absolute_error_baseline.mean(),baseline_minus_main_mae=delta.mean(),
                                    paired_group_ci95_low=lo,paired_group_ci95_high=hi,
                                    folds_favor_main=int((fold_delta>0).sum())))
    pd.DataFrame(audits).to_csv(OUT/'paper_comparison_audit.csv',index=False)
    pd.DataFrame(comparisons).to_csv(OUT/'paper_paired_comparisons.csv',index=False)
    table=pd.read_csv(OUT/'Table3_fair_baselines.csv')
    with pd.ExcelWriter(OUT/'Table3_external_papers.xlsx') as w:
        table.to_excel(w,sheet_name='Fivefold MAE',index=False)
        pd.DataFrame(comparisons).to_excel(w,sheet_name='Paired OOF differences',index=False)
        pd.DataFrame(audits).to_excel(w,sheet_name='OOF audit',index=False)
    doc=Document();doc.add_heading('Table 3. Same-data external baseline comparison',0)
    doc.add_paragraph('Five-fold MAE mean ± sample SD. Lower is better. All methods share the current reconstructed data, target-specific labels and grouped outer folds.')
    t=doc.add_table(rows=1,cols=len(table.columns));t.style='Table Grid'
    for cell,v in zip(t.rows[0].cells,table.columns):cell.text=str(v)
    for _,row in table.iterrows():
        for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
    doc.add_paragraph('Chen et al. (2021), DOI 10.3390/polym13111898: downloaded original model factory, published optimal architecture, Keras 3 torch backend; training-only vocabulary and target scaling; nested early stopping and outer refit. Tg is the original task; the other four targets are method transfers.')
    doc.add_paragraph('Zhang et al. (2024), DOI 10.1021/acs.jctc.3c01385: author-team Polymer-Unit Graph code adapted to current regression tasks. PURS decomposition failed for 2236/8972 structures, which use a whole-repeat-unit MACCS node. Rates vary by target and are detailed in the protocol and conversion audit. This is exploratory adapted author code, not exact paper reproduction.')
    doc.add_paragraph('RF uses a restricted search (leaf sizes 50,200) and is preliminary. Methods have different documented training budgets. Paired bootstrap confidence intervals describe conditional OOF error differences with fixed trained models; they do not cover retraining randomness or corrected multiple testing and are not proof of algorithm-wide significance.')
    doc.save(OUT/'Table3_external_papers.docx')
    lines=['# 外部论文代码迁移与比较','',
           '已使用作者代码在当前数据上训练 Chen RNN；PU-GN 为明确标注的适配版。结果为重新训练，不能称为原论文数值复现。',
           '', '所有可比较结果均核对了测试 row_id、fold、真实标签和有限预测值。',
           '', '差值定义：外部基线 OOF MAE − 完整模型 OOF MAE；正值表示完整模型误差较低。',
           '', '| 方法 | 性质 | 完整模型 OOF MAE | 外部方法 OOF MAE | 差值 | 组 bootstrap 95% CI |',
           '|---|---|---|---|---|---|']
    for r in comparisons:
        lines.append(f"| {r['method']} | {r['target']} | {r['main_oof_mae']:.6g} | {r['baseline_oof_mae']:.6g} | {r['baseline_minus_main_mae']:.6g} | [{r['paired_group_ci95_low']:.6g}, {r['paired_group_ci95_high']:.6g}] |")
    lines += ['', '区间条件于本次已训练模型与固定划分，不含重新训练的随机性；未作多重检验校正，不能据此宣称普遍显著优越。',
              '', '图分解失败比例：Tg 202/557、FFV 1729/7892、Tc 463/866、Density 340/613、Rg 343/614。该图方法的适配结果应独立标注为探索性比较，不能直接替代忠实复现的原方法。',
              '', '未能运行的已核查方法：Kuenneth 仓库未包含预测器；TransTg/Han 未找到匹配的公开实现；Parvez 是分类任务且没有可核实的代码链接。没有为这些论文编造回归成绩。']
    (OUT/'外部论文对比说明.md').write_text('\n'.join(lines)+'\n')
    print(pd.DataFrame(comparisons).to_string(index=False))

if __name__=='__main__':main()
